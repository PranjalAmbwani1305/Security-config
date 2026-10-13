"""Parse a client-supplied file into checklist controls.

The uploaded file is UNTRUSTED input.  Hard limits (size, rows, field length, zip
expansion), an extension allow-list and spreadsheet-formula neutralisation apply.
Nothing in the file is ever executed, and imported controls always start unverified
and land in a *draft* that still needs human review and approval.

Supported: .csv  .json  .xlsx  .txt  .md
"""
import csv
import io
import json
import re
import zipfile

MAX_BYTES = 2 * 1024 * 1024        # 2 MB upload
MAX_CONTROLS = 500
MAX_XLSX_UNCOMPRESSED = 20 * 1024 * 1024
MAX_XLSX_ENTRIES = 200
ALLOWED_EXTENSIONS = (".csv", ".json", ".xlsx", ".txt", ".md")

DEFAULT_CATEGORY = "Imported"
DEFAULT_AUDIT_STEP = "Verify this control is implemented as stated"

# header aliases (lower-cased, spaces/underscores/dashes removed) -> canonical field
_ALIASES = {
    "controlid": "control_id", "itemid": "control_id", "id": "control_id", "ref": "control_id",
    "category": "category", "domain": "category", "section": "category", "area": "category",
    "control": "control", "requirement": "control", "title": "control", "description": "control",
    "controlname": "control", "check": "control",
    "auditstep": "audit_step", "audit": "audit_step", "procedure": "audit_step",
    "verification": "audit_step", "howtoaudit": "audit_step", "auditprocedure": "audit_step",
    "reference": "reference", "standard": "reference", "source": "reference", "benchmark": "reference",
    "rationale": "rationale", "reason": "rationale", "why": "rationale",
}
_LIMITS = {"control_id": 64, "category": 200, "control": 4000, "audit_step": 4000,
           "reference": 300, "rationale": 4000}
_CTRL_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")


class ImportFileError(ValueError):
    """The file cannot be imported; the message is safe to show to the user."""


def _norm_header(h: object) -> str:
    return re.sub(r"[\s_\-]+", "", str(h or "").strip().lower())


def _clean(value: object) -> str:
    s = _CTRL_CHARS.sub("", "" if value is None else str(value)).strip()
    # Neutralise spreadsheet formula injection should the data be exported to CSV/Excel.
    if s and s[0] in "=+-@\t\r":
        s = "'" + s
    return s


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16"):
        try:
            return data.decode(enc)
        except UnicodeError:
            continue
    return data.decode("latin-1")


def _rows_from_records(records: list[dict]) -> list[dict]:
    out = []
    for rec in records:
        row: dict[str, str] = {}
        for k, v in rec.items():
            field = _ALIASES.get(_norm_header(k))
            if field and field not in row and v not in (None, ""):
                row[field] = v
        if row:
            out.append(row)
    return out


def _from_csv(data: bytes) -> list[dict]:
    text = _decode(data)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.DictReader(io.StringIO(text), dialect=dialect)
    if not reader.fieldnames:
        raise ImportFileError("The CSV file has no header row")
    return _rows_from_records(list(reader))


def _from_json(data: bytes) -> list[dict]:
    try:
        doc = json.loads(_decode(data))
    except json.JSONDecodeError as exc:
        raise ImportFileError(f"Invalid JSON: {exc.msg} (line {exc.lineno})") from None
    if isinstance(doc, dict):
        for key in ("checklist", "controls", "items"):
            if isinstance(doc.get(key), list):
                doc = doc[key]
                break
    if not isinstance(doc, list) or not all(isinstance(x, dict) for x in doc):
        raise ImportFileError("JSON must be a list of controls (or an object with a 'checklist'/'controls' list)")
    return _rows_from_records(doc)


def _from_xlsx(data: bytes) -> list[dict]:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            infos = z.infolist()
            if len(infos) > MAX_XLSX_ENTRIES or sum(i.file_size for i in infos) > MAX_XLSX_UNCOMPRESSED:
                raise ImportFileError("The spreadsheet is too large when unpacked")
    except zipfile.BadZipFile:
        raise ImportFileError("This is not a valid .xlsx file") from None
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise ImportFileError("The spreadsheet could not be read") from None
    try:
        ws = wb.worksheets[0]
        rows = ws.iter_rows(values_only=True, max_row=MAX_CONTROLS + 50)
        header = next(rows, None)
        if not header:
            raise ImportFileError("The spreadsheet is empty")
        keys = [str(h or "") for h in header]
        return _rows_from_records([dict(zip(keys, r)) for r in rows])
    finally:
        wb.close()


def _from_text(data: bytes) -> list[dict]:
    """One control per line; bullets/numbering stripped; '# Heading' lines set the category."""
    category, out = "", []
    for line in _decode(data).splitlines():
        line = line.strip()
        if not line:
            continue
        if line.startswith("#"):
            category = line.lstrip("# ").strip()
            continue
        line = re.sub(r"^(?:[-*\u2022]|\d+[.)])\s+", "", line)
        if line:
            row = {"control": line}
            if category:
                row["category"] = category
            out.append(row)
    return out


_PARSERS = {".csv": _from_csv, ".json": _from_json, ".xlsx": _from_xlsx,
            ".txt": _from_text, ".md": _from_text}


def parse_checklist_file(filename: str, data: bytes) -> tuple[list[dict], list[str]]:
    """-> (controls ready for ChecklistControl, warnings).  Raises ImportFileError."""
    name = (filename or "").lower()
    ext = next((e for e in ALLOWED_EXTENSIONS if name.endswith(e)), None)
    if ext is None:
        raise ImportFileError(f"Unsupported file type. Allowed: {', '.join(ALLOWED_EXTENSIONS)}")
    if not data:
        raise ImportFileError("The file is empty")
    if len(data) > MAX_BYTES:
        raise ImportFileError(f"The file is larger than {MAX_BYTES // (1024 * 1024)} MB")

    rows = _PARSERS[ext](data)
    if not rows:
        raise ImportFileError("No controls found. Expected columns such as: control, category, audit_step, reference")
    if len(rows) > MAX_CONTROLS:
        raise ImportFileError(f"Too many controls ({len(rows)}); the limit is {MAX_CONTROLS} per import")

    warnings: list[str] = []
    controls: list[dict] = []
    seen: set[str] = set()
    defaulted_audit = defaulted_cat = 0
    for n, row in enumerate(rows, start=1):
        c = {k: _clean(row.get(k)) for k in _LIMITS}
        if not c["control"]:
            raise ImportFileError(f"Row {n}: the control text is missing")
        for k, limit in _LIMITS.items():
            if len(c[k]) > limit:
                raise ImportFileError(f"Row {n}: '{k}' is longer than {limit} characters")
        if not c["control_id"]:
            c["control_id"] = f"IMP-{n:03d}"
        if c["control_id"] in seen:
            raise ImportFileError(f"Row {n}: duplicate control id '{c['control_id']}'")
        seen.add(c["control_id"])
        if not c["category"]:
            c["category"], defaulted_cat = DEFAULT_CATEGORY, defaulted_cat + 1
        if not c["audit_step"]:
            c["audit_step"], defaulted_audit = DEFAULT_AUDIT_STEP, defaulted_audit + 1
        c["reference_verified"] = False      # a client file is never a verified reference
        controls.append(c)
    if defaulted_cat:
        warnings.append(f"{defaulted_cat} control(s) had no category and were filed under '{DEFAULT_CATEGORY}'")
    if defaulted_audit:
        warnings.append(f"{defaulted_audit} control(s) had no audit step; a generic one was added - please refine it")
    return controls, warnings
